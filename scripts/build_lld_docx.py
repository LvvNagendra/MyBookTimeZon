"""
Build LLD-Appointment-Booking-SaaS.docx from the markdown source + embedded architecture diagram.
Requires: pip install python-docx matplotlib
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
MD_PATH = DOCS / "LLD-Appointment-Booking-SaaS.md"
PNG_PATH = DOCS / "architecture-diagram.png"
DOCX_PATH = DOCS / "LLD-Appointment-Booking-SaaS.docx"
DOCX_FALLBACK = DOCS / "LLD-Appointment-Booking-SaaS-build-output.docx"


def draw_architecture_diagram(out_path: Path) -> None:
    """Reusable stakeholder diagram: multi-vertical tenants + technical stack (Spring Boot + React)."""
    fig, ax = plt.subplots(figsize=(11, 14.5), dpi=150)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 14.2)
    ax.axis("off")
    fig.patch.set_facecolor("#FAFAFA")
    ax.set_facecolor("#FAFAFA")

    def box(x: float, y: float, w: float, h: float, text: str, fc: str = "#E8EAF6", ec: str = "#3949AB") -> None:
        p = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.15",
            facecolor=fc,
            edgecolor=ec,
            linewidth=1.8,
        )
        ax.add_patch(p)
        ax.text(
            x + w / 2,
            y + h / 2,
            text,
            ha="center",
            va="center",
            fontsize=10,
            fontweight="600",
            color="#1A237E",
            wrap=True,
        )

    def cyl(x: float, y: float, w: float, h: float, text: str) -> None:
        """Database cylinder style (simplified as rounded rect)."""
        box(x, y, w, h, text, fc="#E0F2F1", ec="#00695C")

    def arrow(x1: float, y1: float, x2: float, y2: float, dashed: bool = False) -> None:
        style = "--" if dashed else "-"
        arr = FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=14,
            linewidth=1.6,
            color="#424242",
            linestyle=style,
        )
        ax.add_patch(arr)

    # Title
    ax.text(
        5,
        13.65,
        "Appointment Booking SaaS — reusable architecture",
        ha="center",
        fontsize=13,
        fontweight="bold",
        color="#212121",
    )
    ax.text(
        5,
        13.28,
        "One platform for clinics, salons & fitness — shared engine, tenant configuration",
        ha="center",
        fontsize=9,
        color="#424242",
    )

    # --- Multi-vertical reuse (§3.1) ---
    ax.add_patch(
        FancyBboxPatch(
            (0.3, 11.55),
            9.4,
            1.55,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#E8F5E9",
            edgecolor="#2E7D32",
            linewidth=1.3,
        )
    )
    ax.text(0.5, 12.88, "Business verticals (same deployment)", fontsize=9, color="#1B5E20", fontweight="600")
    box(0.5, 12.0, 2.85, 0.78, "Clinics\n(visits, follow-ups)", fc="#C8E6C9", ec="#2E7D32")
    box(3.575, 12.0, 2.85, 0.78, "Salons\n(menu, stylists)", fc="#C8E6C9", ec="#2E7D32")
    box(6.65, 12.0, 2.85, 0.78, "Fitness / PT\n(sessions)", fc="#C8E6C9", ec="#2E7D32")
    box(1.15, 11.62, 7.7, 0.72, "Shared multi-tenant engine  •  clinic_id isolation  •  services / staff / hours per tenant", fc="#A5D6A7", ec="#1B5E20")
    arrow(1.925, 12.0, 2.6, 11.62)
    arrow(5.0, 12.0, 5.0, 11.62)
    arrow(8.075, 12.0, 7.4, 11.62)

    # --- Technical stack (§3.2) ---
    ax.text(5, 11.35, "Technical runtime (Spring Boot + React)", ha="center", fontsize=11, fontweight="bold", color="#37474F")

    # Client tier
    ax.add_patch(
        FancyBboxPatch(
            (0.35, 9.85),
            9.3,
            1.25,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#F5F5F5",
            edgecolor="#9E9E9E",
            linewidth=1.2,
        )
    )
    ax.text(0.55, 10.88, "Client tier", fontsize=9, color="#616161", fontweight="600")
    box(3.2, 9.98, 3.6, 0.92, "React SPA (public + dashboard)")
    arrow(5.0, 11.62, 5.0, 10.9)  # shared engine → React

    # Edge
    ax.add_patch(
        FancyBboxPatch(
            (0.35, 8.25),
            9.3,
            1.25,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#F5F5F5",
            edgecolor="#9E9E9E",
            linewidth=1.2,
        )
    )
    ax.text(0.55, 9.28, "Edge / hosting", fontsize=9, color="#616161", fontweight="600")
    box(2.8, 8.38, 4.4, 0.92, "Vercel / Netlify (static SPA)")

    # API tier
    ax.add_patch(
        FancyBboxPatch(
            (0.35, 6.05),
            9.3,
            1.95,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#FFF8E1",
            edgecolor="#FF8F00",
            linewidth=1.2,
        )
    )
    ax.text(0.55, 7.78, "API tier", fontsize=9, color="#E65100", fontweight="600")
    box(1.0, 6.35, 3.5, 1.05, "Spring Boot API", fc="#FFF3E0", ec="#EF6C00")
    box(5.5, 6.35, 3.5, 1.05, "Spring Security + JWT", fc="#FFF3E0", ec="#EF6C00")

    # Data tier
    ax.add_patch(
        FancyBboxPatch(
            (0.35, 3.95),
            9.3,
            1.85,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#ECEFF1",
            edgecolor="#607D8B",
            linewidth=1.2,
        )
    )
    ax.text(0.55, 5.58, "Data tier", fontsize=9, color="#455A64", fontweight="600")
    cyl(1.2, 4.25, 3.2, 1.15, "PostgreSQL\n(primary DB)")
    cyl(5.6, 4.25, 3.2, 1.15, "Redis (optional)\ncache / locks")

    # Integrations
    ax.add_patch(
        FancyBboxPatch(
            (0.35, 1.55),
            9.3,
            2.15,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor="#F3E5F5",
            edgecolor="#7B1FA2",
            linewidth=1.2,
        )
    )
    ax.text(0.55, 3.48, "Integrations", fontsize=9, color="#6A1B9A", fontweight="600")
    box(0.55, 2.62, 2.7, 0.78, "Razorpay /\nStripe", fc="#F3E5F5", ec="#7B1FA2")
    box(3.65, 2.62, 2.7, 0.78, "Twilio SMS /\nWhatsApp", fc="#F3E5F5", ec="#7B1FA2")
    box(6.75, 2.62, 2.7, 0.78, "Email\nprovider", fc="#F3E5F5", ec="#7B1FA2")

    cx = 5.0
    arrow(cx, 9.98, cx, 9.5)
    arrow(cx, 8.38, cx, 8.0)
    arrow(2.75, 6.88, 5.45, 6.88)

    arrow(2.25, 6.35, 2.7, 5.4)
    arrow(7.75, 6.35, 7.3, 5.4, dashed=True)

    arrow(2.7, 4.25, 1.9, 3.45)
    arrow(5.0, 4.25, 5.0, 3.45)
    arrow(7.3, 4.25, 8.1, 3.45)

    ax.text(
        5.0,
        1.12,
        "Solid: primary flows  |  Dashed: optional Redis  |  Diagram reusable for proposals & onboarding",
        ha="center",
        fontsize=8,
        color="#616161",
    )

    plt.tight_layout()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_path, bbox_inches="tight", facecolor=fig.patch.get_facecolor())
    plt.close(fig)


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table_from_md(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        parts = [c.strip() for c in line.strip().strip("|").split("|")]
        if set(parts) <= {"", "-", ":"} or all(re.match(r"^:?-+:?$", p) for p in parts if p):
            continue
        rows.append(parts)
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    for i, row_data in enumerate(rows):
        for j in range(ncol):
            cell = table.rows[i].cells[j]
            text = row_data[j] if j < len(row_data) else ""
            cell.text = text
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
            if i == 0:
                set_cell_shading(cell, "D9E2F3")
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(10)
    doc.add_paragraph()


def parse_md_to_docx(md_text: str, image_path: Path) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_text.replace("\r\n", "\n").split("\n")
    first_heading_done = False
    i = 0
    in_fence = False
    fence_lang = ""
    fence_lines: list[str] = []
    table_buf: list[str] = []

    def flush_table() -> None:
        nonlocal table_buf
        if table_buf:
            add_table_from_md(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i]

        if in_fence:
            if line.strip().startswith("```"):
                if fence_lang == "mermaid":
                    pass  # diagram replaced by PNG in section 3
                else:
                    add_code_block(doc, fence_lines)
                in_fence = False
                fence_lang = ""
                fence_lines = []
                i += 1
                continue
            fence_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            flush_table()
            in_fence = True
            fence_lang = line.strip()[3:].strip().lower()
            fence_lines = []
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        else:
            flush_table()

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith("#"):
            flush_table()
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            hl = min(max(level, 1), 9)
            h = doc.add_heading(text, level=hl)
            if not first_heading_done and level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_heading_done = True
            elif level == 2 and doc.paragraphs[-1].text.startswith("Appointment Booking"):
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Insert architecture image right after section 3 heading
            if text.lower().startswith("3. logical architecture"):
                doc.add_paragraph(
                    "The diagram below combines multi-vertical reuse (clinics, salons, fitness) "
                    "with the technical runtime: React, edge hosting, Spring Boot, data stores, and integrations."
                )
                if image_path.exists():
                    doc.add_picture(str(image_path), width=Inches(6.5))
                    last_p = doc.paragraphs[-1]
                    last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    doc.add_paragraph("[Architecture image not found — run script to generate PNG.]")
            i += 1
            continue

        if line.strip() == "":
            doc.add_paragraph()
            i += 1
            continue

        # Normal paragraph (strip simple md bold)
        t = line
        t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
        t = re.sub(r"`([^`]+)`", r"\1", t)
        p = doc.add_paragraph(t)
        i += 1

    flush_table()
    return doc


def main() -> int:
    if not MD_PATH.exists():
        print(f"Missing markdown: {MD_PATH}", file=sys.stderr)
        return 1
    print("Drawing architecture diagram...")
    draw_architecture_diagram(PNG_PATH)
    print(f"Wrote {PNG_PATH}")
    md = MD_PATH.read_text(encoding="utf-8")
    # Remove mermaid block from text flow when parsing — parser skips mermaid body
    doc = parse_md_to_docx(md, PNG_PATH)
    try:
        doc.save(DOCX_PATH)
        out = DOCX_PATH
    except PermissionError:
        doc.save(DOCX_FALLBACK)
        out = DOCX_FALLBACK
        print(
            f"Could not overwrite {DOCX_PATH} (file open?). Wrote {DOCX_FALLBACK} instead.",
            file=sys.stderr,
        )
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
